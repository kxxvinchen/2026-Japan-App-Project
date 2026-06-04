from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# (Day, Date, City, Hotel, Performances, Other Activities)
rows = [
    ("Day 1", "Tue, Jun 9, 2026", "San Francisco / In-flight",
     "Overnight en route (in-flight)",
     "—",
     "Board coaches at Saratoga HS, transfer to San Francisco International Airport (SFO). Check in luggage and instruments, clear security, depart for Tokyo. Overnight en route."),
    ("Day 2", "Wed, Jun 10, 2026", "Tokyo",
     "Grand Nikko Tokyo Daiba (Students/Staff/Chaperones); Tokyo Metropolitan Hotel Ikebukuro (Shadow Tour)",
     "—",
     "Arrive Tokyo (Haneda HND or Narita NRT), clear Immigration/Customs, meet local Tour Managers. Hotel check-in. Dinner on own – Aqua City waterfront mall (SWE/SSO/Choir) or Ikebukuro (Shadow). Free evening until curfew."),
    ("Day 3", "Thu, Jun 11, 2026", "Tokyo",
     "Grand Nikko Tokyo Daiba (Students/Staff/Chaperones); Tokyo Metropolitan Hotel Ikebukuro (Shadow Tour)",
     "—  (rehearsals only: Nishi Arai Cultural Hall for SWE/SSO, Katsushika Symphony Hills for Choir, Iwakura HS for Percussion)",
     "SWE/SSO/Choir: Tokyo Skytree and Asakusa (Sensō-ji Temple). Shadow: Nijubashi Bridge, Tokyo Station, Ginza, Meiji Shrine, Harajuku (Takeshita St, Omotesando)."),
    ("Day 4", "Fri, Jun 12, 2026", "Tokyo",
     "Grand Nikko Tokyo Daiba (Students/Staff/Chaperones); Tokyo Metropolitan Hotel Ikebukuro (Shadow Tour)",
     "Choir: performance at Toka Hall — Friendship Concert with Toko Jr. & Sr. High School.",
     "Harajuku – Takeshita Street and Meiji Shrine. Shadow: Skytree, Asakusa, Kappabashi (Kitchen Town). SWE/SSO cultural exchange at Ichikawa Jr. & Sr. HS. Group dinner in Shinjuku/Shibuya."),
    ("Day 5", "Sat, Jun 13, 2026", "Tokyo",
     "Grand Nikko Tokyo Daiba (Students/Staff/Chaperones); Tokyo Metropolitan Hotel Ikebukuro (Shadow Tour)",
     "SWE/SSO: performance at Beethoven Hall, Musashino Academia Musicae — Friendship Concert with Iwakura HS Wind Orchestra. Choir: performance at Ikebukuro Global Ring.",
     "Shadow & Choir: day trip to Kawagoe (\"Little Edo\"). Shadow attends the concerts."),
    ("Day 6", "Sun, Jun 14, 2026", "Tokyo → Fuji Lake Kawaguchi",
     "Hotel Mt. Fuji (Band & Choir) / Highland Resort Hotel and Spa (Orchestra); Fuji Premium Resort (Choir Shadow) / The Prince Hakone Lake Ashi (Band/Orchestra Shadow)",
     "SWE/SSO: performance at Stellar Theater — Mt. Fuji Music Festival. Choir: performance at Shopping Center Bell with local choirs.",
     "Check out, transfer toward Mt. Fuji area (~2.5 hrs). Stops at Oshino Hakkai and Mt. Fuji 5th Station. Group dinner at hotel."),
    ("Day 7", "Mon, Jun 15, 2026", "Fuji Lake Kawaguchi → Kyoto",
     "Miyako Hotel Kyoto Hachijo (Students/Staff/Chaperones); Hotel Nikko Princess Kyoto (Shadow Tour)",
     "Choir: performance at the Organ/Concert Hall of Kawaguchiko Music Forest Museum.",
     "SWE/SSO: Oshino Hakkai and Mt. Fuji 5th Station. Band/Orchestra Shadow: free time at Hakone-En. Transfer to Kyoto via Shinkansen (bullet train). Hotel check-in."),
    ("Day 8", "Tue, Jun 16, 2026", "Kyoto",
     "Miyako Hotel Kyoto Hachijo (Students/Staff/Chaperones); Hotel Nikko Princess Kyoto (Shadow Tour)",
     "—",
     "Free half-day in Kyoto – optional Kinkaku-ji (Golden Pavilion), Kiyomizu-dera, Nijō Castle, Sanjūsangen-dō, Gion district. Afternoon excursion to Arashiyama and the Bamboo Grove. Group dinner in Arashiyama."),
    ("Day 9", "Wed, Jun 17, 2026", "Kyoto → Kobe",
     "Kobe Bay Sheraton Hotel & Towers (Band & Orchestra) / Kobe Plaza Hotel (Choir); Grand Nikko Awaji (Shadow Tour)",
     "SWE/SSO: performance at Akashi Civic Hall — Friendship Concert with Takigawa Daini HS.",
     "Transfer to Kobe (~2 hrs). Shadow: Fushimi Inari, Fushimi Sake District, Arashiyama bamboo forest. Hotel check-in."),
    ("Day 10", "Thu, Jun 18, 2026", "Kobe",
     "Kobe Bay Sheraton Hotel & Towers (Band & Orchestra) / Kobe Plaza Hotel (Choir); Grand Nikko Awaji (Shadow Tour)",
     "Choir: Friendship Concert at Kobe High School Auditorium.",
     "Relaxing morning; optional activities (Zen experience, Awaji Hello Kitty Apple Land, etc.). Kobe city highlights tour – Chinatown, Kobe Harborland, Meriken Park, Port Tower. Dinner on own."),
    ("Day 11", "Fri, Jun 19, 2026", "Kobe → Hiroshima",
     "Grand Prince Hotel Hiroshima (Students/Staff/Chaperones); Rihga Royal Hotel Hiroshima (Shadow Tour)",
     "Choir: short performance in Hiroshima Peace Memorial Park.",
     "Check out, transfer to Hiroshima via Shinkansen (~1 hr 10 min). Visit Hiroshima Peace Memorial Park and the Atomic Bomb Dome (UNESCO World Heritage Site). Hotel check-in, group dinner."),
    ("Day 12", "Sat, Jun 20, 2026", "Hiroshima",
     "Grand Prince Hotel Hiroshima (Students/Staff/Chaperones); Rihga Royal Hotel Hiroshima (Shadow Tour)",
     "SWE/SSO: Peace Friendship Concert at Phoenix Concert Hall with Hijiyama Girls Jr. & Sr. HS Wind Orchestra (reunion concert). Choir: performance at Jogakuin HS Gaine's Chapel Hall.",
     "Choir: visit Jogakuin Girls HS for classes and cultural exchange. SWE/SSO pack instruments after the concert. Dinner in downtown Hiroshima."),
    ("Day 13", "Sun, Jun 21, 2026", "Hiroshima (Miyajima excursion)",
     "Grand Prince Hotel Hiroshima (all travelers, incl. Shadow Tour)",
     "—",
     "SWE/SSO/Choir: cultural exchange at Hijiyama HS (cultural/musical performances, paper crane making). Excursion to Miyajima island – floating torii gate of Itsukushima Shrine. Farewell Dinner Dance Party at the hotel."),
    ("Day 14", "Mon, Jun 22, 2026", "Hiroshima → Tokyo",
     "Grand Nikko Tokyo Daiba (all travelers, incl. Shadow Tour)",
     "—",
     "Check out, transfer to Hiroshima Station, board Shinkansen to Tokyo (~4 hrs). Tokyo Dome professional baseball game. Hotel check-in; last evening in Japan."),
    ("Day 15", "Tue, Jun 23, 2026", "Tokyo → San Francisco",
     "— (departure day)",
     "—",
     "Breakfast, check out, transfer to airport. Depart Narita (NRT) or Haneda (HND) for home. Arrive back at SFO; collect luggage and instruments."),
]

wb = Workbook()
ws = wb.active
ws.title = "Japan Tour Itinerary"

title = "2026 Japan Performance & Cultural Exchange Tour — Saratoga HS & Lynbrook HS"
subtitle = "June 9–23, 2026  •  Including Mt. Fuji Music Festival  •  Arranged by BMC Tours"

ws["A1"] = title
ws.merge_cells("A1:F1")
ws["A1"].font = Font(name="Arial", size=14, bold=True, color="FFFFFF")
ws["A1"].fill = PatternFill("solid", start_color="1F3864")
ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
ws.row_dimensions[1].height = 30

ws["A2"] = subtitle
ws.merge_cells("A2:F2")
ws["A2"].font = Font(name="Arial", size=10, italic=True, color="1F3864")
ws["A2"].alignment = Alignment(horizontal="center", vertical="center")
ws.row_dimensions[2].height = 20

headers = ["Day", "Date", "City", "Hotel", "Performances", "Other Activities"]
ws.append(headers)
hdr_fill = PatternFill("solid", start_color="2F5496")
for col, _ in enumerate(headers, start=1):
    c = ws.cell(row=3, column=col)
    c.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    c.fill = hdr_fill
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
ws.row_dimensions[3].height = 22

thin = Side(style="thin", color="BFBFBF")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
perf_font_color = "1F3864"

for i, row in enumerate(rows):
    ws.append(row)
    r = 4 + i
    fill = "F2F2F2" if i % 2 else "FFFFFF"
    for col in range(1, 7):
        c = ws.cell(row=r, column=col)
        c.font = Font(name="Arial", size=10)
        c.fill = PatternFill("solid", start_color=fill)
        c.border = border
        wrap = col in (3, 4, 5, 6)
        c.alignment = Alignment(horizontal="left", vertical="top", wrap_text=wrap)
    ws.cell(row=r, column=1).font = Font(name="Arial", size=10, bold=True)
    ws.cell(row=r, column=1).alignment = Alignment(horizontal="center", vertical="top")
    ws.cell(row=r, column=2).alignment = Alignment(horizontal="left", vertical="top")
    # Emphasize real performances
    if not rows[i][4].startswith("—"):
        ws.cell(row=r, column=5).font = Font(name="Arial", size=10, bold=True, color=perf_font_color)

widths = {"A": 8, "B": 18, "C": 28, "D": 38, "E": 48, "F": 55}
for col, w in widths.items():
    ws.column_dimensions[col].width = w

for i in range(len(rows)):
    ws.row_dimensions[4 + i].height = 105

ws.freeze_panes = "A4"
ws.sheet_view.showGridLines = False

wb.save("/sessions/laughing-determined-cori/mnt/outputs/Japan_Tour_Itinerary_2026.xlsx")
print("saved")
